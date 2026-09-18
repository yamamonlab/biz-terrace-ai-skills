#!/usr/bin/env python3
"""Basic PPTX geometry inspection.

Checks shapes outside slide bounds and significant overlaps between text-containing shapes.
This is a geometric check, not a substitute for visual rendering inspection.
"""
from __future__ import annotations
import argparse, json


def intersection(a,b):
    x1=max(a[0],b[0]); y1=max(a[1],b[1]); x2=min(a[2],b[2]); y2=min(a[3],b[3])
    if x2<=x1 or y2<=y1: return 0
    return (x2-x1)*(y2-y1)


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument("pptx")
    args=ap.parse_args()
    try:
        from pptx import Presentation
    except ImportError:
        raise SystemExit("python-pptx is required")
    prs=Presentation(args.pptx)
    W,H=prs.slide_width,prs.slide_height
    issues=[]
    for si,slide in enumerate(prs.slides,1):
        boxes=[]
        for j,sh in enumerate(slide.shapes):
            x,y,w,h=sh.left,sh.top,sh.width,sh.height
            if x<0 or y<0 or x+w>W or y+h>H:
                issues.append({"severity":"ERROR","code":"R01","slide":si,"shape":j,"message":"shape outside slide bounds"})
            text=""
            if hasattr(sh,"text"):
                text=(sh.text or "").strip()
            if text:
                boxes.append((j,(x,y,x+w,y+h),w*h,text[:40]))
        for a in range(len(boxes)):
            for b in range(a+1,len(boxes)):
                ia=intersection(boxes[a][1],boxes[b][1])
                if ia and ia/min(boxes[a][2],boxes[b][2]) > 0.20:
                    issues.append({"severity":"WARN","code":"R02","slide":si,"shapes":[boxes[a][0],boxes[b][0]],"message":"text shapes overlap >20% of smaller box"})
    print(json.dumps(issues,ensure_ascii=False,indent=2))
    raise SystemExit(1 if any(x["severity"]=="ERROR" for x in issues) else 0)

if __name__ == "__main__":
    main()
